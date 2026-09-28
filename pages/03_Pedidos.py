import streamlit as st
import pandas as pd
import datetime
import io
from modules.auth import requer_permissao
from modules.google_sheets import ler_df, escrever_df, append_linha, ler_categorias, atualizar_linha

usuario = requer_permissao("pedidos")

st.title("📋 Solicitações de Compra")

df_produtos = ler_df("produtos")
df_pedidos  = ler_df("pedidos")
df_itens    = ler_df("itens_pedido")

perfil = usuario["perfil"]
unidades_acesso = usuario["unidades_acesso"]
if unidades_acesso == "todos":
    unidades_disponiveis = ler_df("unidades")["nome"].tolist()
else:
    unidades_disponiveis = [u.strip() for u in str(unidades_acesso).split(",")]


def _safe_int(v):
    try:
        return int(float(v)) if str(v).strip() not in ("", "nan") else 0
    except Exception:
        return 0


def _next_id(df, col="id"):
    if df.empty or col not in df.columns:
        return 1
    try:
        return int(df[col].apply(_safe_int).max()) + 1
    except Exception:
        return 1


# ── Produtos ativos organizados por categoria (usado em Nova Solicitação e Editar) ──
produtos_ativos = (
    df_produtos[df_produtos["ativo"].astype(str).str.upper().isin(["TRUE", "1", "SIM"])].copy()
    if not df_produtos.empty else pd.DataFrame()
)
if not produtos_ativos.empty:
    if "categoria" not in produtos_ativos.columns:
        produtos_ativos["categoria"] = ""
    produtos_ativos["categoria"] = produtos_ativos["categoria"].fillna("").astype(str).str.strip()
    produtos_ativos = produtos_ativos.sort_values(
        "descricao", key=lambda x: x.str.lower(), ignore_index=True
    )
    _ordem_cats    = ler_categorias() + [""]
    _cats_presentes = [c for c in _ordem_cats if c in produtos_ativos["categoria"].values]
else:
    _cats_presentes = []

CSS = """
<style>
.prod-row {
    display: flex;
    align-items: baseline;
    width: 100%;
    gap: 0;
}
.prod-nome {
    white-space: nowrap;
    font-weight: 600;
    font-size: 14px;
}
.prod-unidade {
    white-space: nowrap;
    font-size: 12px;
    color: #888;
    margin-left: 4px;
}
.prod-dots {
    flex: 1;
    border-bottom: 1.5px dotted #bbb;
    margin: 0 8px 4px 8px;
    min-width: 20px;
}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


def _renderizar_form_produtos(sufixo: str, qtds_pre: dict, mostrar_salvar: bool = False) -> dict:
    """Renderiza o formulário de produtos com quantidades pré-preenchidas.
    Retorna dict {produto_id: quantidade} com os itens preenchidos."""
    itens = {}
    for cat in _cats_presentes:
        grupo = produtos_ativos[produtos_ativos["categoria"] == cat]
        if grupo.empty:
            continue
        cat_label = cat if cat else "Sem categoria"
        with st.expander(f"📦 {cat_label}", expanded=True):
            for _, prod in grupo.iterrows():
                pid = str(prod["id"])
                col1, col2 = st.columns([4, 1])
                with col1:
                    st.markdown(
                        f'<div class="prod-row">'
                        f'<span class="prod-nome">{prod["descricao"]}</span>'
                        f'<span class="prod-unidade">— {prod["unidade_base"]}</span>'
                        f'<span class="prod-dots"></span>'
                        f'</div>',
                        unsafe_allow_html=True,
                    )
                    if prod.get("observacao"):
                        st.caption(prod["observacao"])
                with col2:
                    qtd = st.number_input(
                        "Qtd",
                        min_value=0.0,
                        step=0.5,
                        value=float(qtds_pre.get(pid, 0.0)),
                        key=f"qtd_{pid}_{sufixo}",
                        label_visibility="collapsed",
                    )
                    if qtd > 0:
                        itens[prod["id"]] = qtd

            if mostrar_salvar:
                cat_key = cat.replace(" ", "_") if cat else "sem_cat"
                if st.button(
                    f"💾 Salvar rascunho — {cat_label}",
                    key=f"salvar_rascunho_{cat_key}_{sufixo}",
                    use_container_width=True,
                ):
                    for _, p in grupo.iterrows():
                        pk = str(p["id"])
                        v = st.session_state.get(f"qtd_{pk}_{sufixo}", 0.0)
                        if v and float(v) > 0:
                            st.session_state["rascunho_pedido"][pk] = float(v)
                        else:
                            st.session_state["rascunho_pedido"].pop(pk, None)
                    st.toast(f"Rascunho da categoria '{cat_label}' salvo!", icon="💾")
    return itens


# ── Session state ────────────────────────────────────────────────────────────
if "pedido_form_v" not in st.session_state:
    st.session_state["pedido_form_v"] = 0
if "editando_pedido" not in st.session_state:
    st.session_state["editando_pedido"] = None
if "rascunho_pedido" not in st.session_state:
    st.session_state["rascunho_pedido"] = {}
if "confirmar_bloquear" not in st.session_state:
    st.session_state["confirmar_bloquear"] = None
if "confirmar_cancelar" not in st.session_state:
    st.session_state["confirmar_cancelar"] = None

tab_novo, tab_abertos, tab_importar = st.tabs(["Nova Solicitação", "Solicitações Abertas", "📥 Importar Planilha"])

# ── TAB: NOVA SOLICITAÇÃO ────────────────────────────────────────────────────
with tab_novo:
    unidade = st.selectbox("Unidade", unidades_disponiveis)

    if produtos_ativos.empty:
        st.warning("Nenhum produto ativo cadastrado.")
    else:
        st.markdown("Preencha a quantidade desejada. Use **Salvar rascunho** por categoria para não perder o progresso.")
        fv = st.session_state["pedido_form_v"]
        rascunho = st.session_state["rascunho_pedido"]
        itens_pedido = _renderizar_form_produtos(sufixo=f"novo_{fv}", qtds_pre=rascunho, mostrar_salvar=True)

        if st.button("Enviar Solicitação", type="primary"):
            if not itens_pedido:
                st.error("Selecione ao menos um produto com quantidade.")
            else:
                try:
                    novo_id = _next_id(df_pedidos)
                    item_id = _next_id(df_itens)
                    append_linha("pedidos", [
                        novo_id, unidade, "aberto",
                        usuario["nome"], datetime.datetime.now().isoformat(), "", "",
                    ])
                    for prod_id, qtd in itens_pedido.items():
                        append_linha("itens_pedido", [item_id, novo_id, prod_id, qtd])
                        item_id += 1
                    st.session_state["pedido_form_v"] += 1
                    st.session_state["rascunho_pedido"] = {}
                    st.cache_data.clear()
                    st.success(f"Solicitação #{novo_id} enviada para {unidade}!")
                    st.rerun()
                except Exception as e:
                    st.error(f"Erro ao salvar solicitação: {e}")

# ── TAB: SOLICITAÇÕES ABERTAS ────────────────────────────────────────────────
with tab_abertos:
    if df_pedidos.empty:
        st.info("Nenhuma solicitação aberta.")
    else:
        pedidos_abertos = df_pedidos[df_pedidos["status"] == "aberto"].copy()
        if unidades_acesso != "todos":
            pedidos_abertos = pedidos_abertos[pedidos_abertos["unidade"].isin(unidades_disponiveis)]

        if pedidos_abertos.empty:
            st.info("Nenhuma solicitação aberta para sua(s) unidade(s).")
        else:
            for _, ped in pedidos_abertos.iterrows():
                ped_id     = _safe_int(ped["id"])
                criado_por = str(ped.get("criado_por", "") or "")
                data_label = str(ped.get("data_criacao", ""))[:10] or "—"
                editado_por = str(ped.get("editado_por", "") or "")
                data_edicao = str(ped.get("data_edicao", "") or "")

                pode_editar   = perfil == "admin" or criado_por == usuario["nome"]
                pode_bloquear = perfil in ["admin", "comprador"]
                pode_cancelar = perfil == "admin" or criado_por == usuario["nome"]

                with st.expander(f"#{ped_id} — {ped['unidade']} — {data_label} — por {criado_por}"):
                    itens = (
                        df_itens[df_itens["pedido_id"].apply(_safe_int) == ped_id]
                        if not df_itens.empty else pd.DataFrame()
                    )

                    # ── Modo edição ──────────────────────────────────────────
                    if st.session_state["editando_pedido"] == ped_id:
                        st.caption("✏️ Editando solicitação — ajuste as quantidades e salve.")

                        qtds_pre = {}
                        if not itens.empty:
                            for _, it in itens.iterrows():
                                try:
                                    qtds_pre[str(it["produto_id"])] = float(it["quantidade"])
                                except Exception:
                                    pass

                        novos_itens = _renderizar_form_produtos(
                            sufixo=f"edit_{ped_id}", qtds_pre=qtds_pre
                        )

                        col_s, col_c = st.columns(2)
                        with col_s:
                            if st.button("💾 Salvar Alterações", key=f"salvar_edit_{ped_id}", type="primary", use_container_width=True):
                                if not novos_itens:
                                    st.error("Selecione ao menos um produto com quantidade.")
                                else:
                                    try:
                                        df_itens_reload = ler_df("itens_pedido")
                                        df_sem_este = (
                                            df_itens_reload[df_itens_reload["pedido_id"].apply(_safe_int) != ped_id]
                                            if not df_itens_reload.empty else pd.DataFrame()
                                        )
                                        escrever_df("itens_pedido", df_sem_este)
                                        item_id = _next_id(df_sem_este)
                                        for prod_id, qtd in novos_itens.items():
                                            append_linha("itens_pedido", [item_id, ped_id, prod_id, qtd])
                                            item_id += 1
                                        atualizar_linha("pedidos", ped["id"], {
                                            "editado_por": usuario["nome"],
                                            "data_edicao": datetime.datetime.now().isoformat(),
                                        })
                                        st.session_state["editando_pedido"] = None
                                        st.cache_data.clear()
                                        st.success("Solicitação atualizada!")
                                        st.rerun()
                                    except Exception as e:
                                        st.error(f"Erro ao salvar: {e}")
                        with col_c:
                            if st.button("✖ Cancelar", key=f"cancelar_edit_{ped_id}", use_container_width=True):
                                st.session_state["editando_pedido"] = None
                                st.rerun()

                    # ── Modo visualização ────────────────────────────────────
                    else:
                        if not itens.empty and not df_produtos.empty:
                            itens_display = itens.merge(
                                df_produtos[["id", "descricao", "unidade_base"]],
                                left_on="produto_id", right_on="id", how="left"
                            )[["descricao", "unidade_base", "quantidade"]]
                            st.dataframe(itens_display, use_container_width=True, hide_index=True)

                        if editado_por:
                            st.caption(f"✏️ Última edição por **{editado_por}** em {data_edicao[:10] if data_edicao else '—'}")

                        n_botoes = sum([pode_editar, pode_bloquear, pode_cancelar])
                        if n_botoes:
                            cols = st.columns(n_botoes)
                            ci = 0

                            if pode_editar:
                                with cols[ci]:
                                    if st.button("✏️ Editar", key=f"editar_{ped_id}", use_container_width=True):
                                        st.session_state["editando_pedido"] = ped_id
                                        st.session_state["confirmar_bloquear"] = None
                                        st.session_state["confirmar_cancelar"] = None
                                        st.rerun()
                                ci += 1

                            if pode_bloquear:
                                with cols[ci]:
                                    if st.session_state["confirmar_bloquear"] == ped_id:
                                        st.warning("Confirmar bloqueio desta solicitação?")
                                        cb1, cb2 = st.columns(2)
                                        with cb1:
                                            if st.button("✅ Confirmar", key=f"conf_bloquear_{ped_id}", type="primary", use_container_width=True):
                                                atualizar_linha("pedidos", ped["id"], {
                                                    "status": "bloqueado",
                                                    "data_bloqueio": datetime.datetime.now().isoformat(),
                                                })
                                                st.session_state["confirmar_bloquear"] = None
                                                st.success("Pedido bloqueado!")
                                                st.cache_data.clear()
                                                st.rerun()
                                        with cb2:
                                            if st.button("✖ Voltar", key=f"canc_bloquear_{ped_id}", use_container_width=True):
                                                st.session_state["confirmar_bloquear"] = None
                                                st.rerun()
                                    else:
                                        if st.button("🔒 Bloquear (consolidar)", key=f"bloquear_{ped_id}", use_container_width=True):
                                            st.session_state["confirmar_bloquear"] = ped_id
                                            st.session_state["confirmar_cancelar"] = None
                                            st.rerun()
                                ci += 1

                            if pode_cancelar:
                                with cols[ci]:
                                    if st.session_state["confirmar_cancelar"] == ped_id:
                                        st.warning("Confirmar cancelamento desta solicitação?")
                                        cc1, cc2 = st.columns(2)
                                        with cc1:
                                            if st.button("✅ Confirmar", key=f"conf_cancelar_{ped_id}", type="primary", use_container_width=True):
                                                atualizar_linha("pedidos", ped["id"], {
                                                    "status": "cancelado",
                                                    "editado_por": usuario["nome"],
                                                    "data_edicao": datetime.datetime.now().isoformat(),
                                                })
                                                st.session_state["confirmar_cancelar"] = None
                                                st.success("Solicitação cancelada.")
                                                st.cache_data.clear()
                                                st.rerun()
                                        with cc2:
                                            if st.button("✖ Voltar", key=f"canc_cancelar_{ped_id}", use_container_width=True):
                                                st.session_state["confirmar_cancelar"] = None
                                                st.rerun()
                                    else:
                                        if st.button("❌ Cancelar Solicitação", key=f"cancelar_{ped_id}", use_container_width=True):
                                            st.session_state["confirmar_cancelar"] = ped_id
                                            st.session_state["confirmar_bloquear"] = None
                                            st.rerun()
                                ci += 1

# ── TAB: IMPORTAR PLANILHA ────────────────────────────────────────────────────
with tab_importar:
    st.markdown(
        "Baixe o modelo, preencha as quantidades por estabelecimento e faça o upload "
        "para criar todas as solicitações de uma vez."
    )

    # ── Gerador de template ───────────────────────────────────────────────────
    def _gerar_template_xlsx() -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Solicitação"

        header_fixo = ["codigo", "descricao", "apresentacao", "unidade_base"]
        header = header_fixo + unidades_disponiveis
        ws.append(header)

        # Estilo cabeçalho fixo
        fill_cinza  = PatternFill("solid", fgColor="D9D9D9")
        fill_azul   = PatternFill("solid", fgColor="BDD7EE")
        bold        = Font(bold=True)
        centro      = Alignment(horizontal="center", vertical="center")
        borda_thin  = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"),  bottom=Side(style="thin"),
        )

        for col_idx, _ in enumerate(header, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = bold
            cell.alignment = centro
            cell.border = borda_thin
            cell.fill = fill_azul if col_idx > len(header_fixo) else fill_cinza

        # Estilos de linha de categoria
        fill_cat    = PatternFill("solid", fgColor="F2F2F2")
        bold_italic = Font(bold=True, italic=True, size=10)

        # Agrupa por categoria (somente produtos ativos — inativos já excluídos)
        col_cat = "categoria" if "categoria" in produtos_ativos.columns else None
        if col_cat:
            categorias_ordem = produtos_ativos[col_cat].fillna("Sem categoria").unique().tolist()
        else:
            categorias_ordem = ["Produtos"]

        for cat in categorias_ordem:
            if col_cat:
                grupo = produtos_ativos[produtos_ativos[col_cat].fillna("Sem categoria") == cat]
            else:
                grupo = produtos_ativos

            if grupo.empty:
                continue

            # Linha de título da categoria
            ws.append([cat] + [""] * (len(header) - 1))
            row_cat = ws.max_row
            for col_idx in range(1, len(header) + 1):
                cell = ws.cell(row=row_cat, column=col_idx)
                cell.fill = fill_cat
                cell.font = bold_italic
                cell.border = borda_thin

            # Produtos da categoria
            for _, prod in grupo.iterrows():
                row_data = [
                    str(prod.get("codigo", "") or ""),
                    str(prod.get("descricao", "") or ""),
                    str(prod.get("apresentacao", "") or ""),
                    str(prod.get("unidade_base", "") or ""),
                ] + [""] * len(unidades_disponiveis)
                ws.append(row_data)

        # Larguras de coluna
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 12
        for i in range(len(unidades_disponiveis)):
            ws.column_dimensions[get_column_letter(5 + i)].width = 16

        # Travar cabeçalho e colunas fixas
        ws.freeze_panes = "E2"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    col_dl_t, _ = st.columns([2, 4])
    with col_dl_t:
        if not produtos_ativos.empty:
            st.download_button(
                "⬇️ Baixar modelo Excel",
                data=_gerar_template_xlsx(),
                file_name=f"modelo_solicitacao_{datetime.date.today()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        else:
            st.warning("Nenhum produto ativo para gerar o modelo.")

    st.markdown("---")
    st.markdown("**Fazer upload da planilha preenchida:**")

    if "import_ok" in st.session_state and st.session_state["import_ok"]:
        st.success(st.session_state["import_ok"])
        if st.button("📥 Importar outra planilha", use_container_width=True):
            del st.session_state["import_ok"]
            st.session_state["upload_v"] = st.session_state.get("upload_v", 0) + 1
            st.rerun()
        st.stop()

    if "upload_v" not in st.session_state:
        st.session_state["upload_v"] = 0

    arq = st.file_uploader(
        "Selecione o arquivo Excel (.xlsx)",
        type=["xlsx"],
        key=f"upload_planilha_{st.session_state['upload_v']}",
    )

    if arq:
        try:
            df_upload = pd.read_excel(arq, dtype=str)
        except Exception as e:
            st.error(f"Erro ao ler o arquivo: {e}")
            st.stop()

        # Identificar colunas de unidades (tudo após as 4 colunas fixas)
        colunas_fixas = ["codigo", "descricao", "apresentacao", "unidade_base"]
        colunas_unid_excel = [c for c in df_upload.columns if c not in colunas_fixas]

        if not colunas_unid_excel:
            st.error("A planilha não contém colunas de unidades. Verifique se usou o modelo correto.")
        else:
            # Mapear descricao → produto_id (normalizado)
            mapa_prod = {}
            if not produtos_ativos.empty:
                for _, p in produtos_ativos.iterrows():
                    chave = str(p["descricao"]).strip().lower()
                    mapa_prod[chave] = p["id"]
                    if str(p.get("codigo", "") or "").strip():
                        mapa_prod[str(p["codigo"]).strip().lower()] = p["id"]

            # Coletar itens por unidade
            itens_por_unidade = {u: {} for u in colunas_unid_excel}
            nao_encontrados = set()

            for _, row in df_upload.iterrows():
                descricao = str(row.get("descricao", "") or "").strip()
                codigo    = str(row.get("codigo", "") or "").strip()
                if not descricao or descricao.lower() == "nan":
                    continue

                prod_id = (
                    mapa_prod.get(codigo.lower()) or
                    mapa_prod.get(descricao.lower())
                )

                if prod_id is None:
                    nao_encontrados.add(descricao)
                    continue

                for unid in colunas_unid_excel:
                    val = str(row.get(unid, "") or "").strip().replace(",", ".")
                    if val and val not in ("", "nan", "None"):
                        try:
                            qtd = float(val)
                            if qtd > 0:
                                itens_por_unidade[unid][prod_id] = qtd
                        except ValueError:
                            pass

            if nao_encontrados:
                st.warning(
                    f"⚠️ {len(nao_encontrados)} produto(s) da planilha não foram encontrados no cadastro "
                    f"e serão ignorados: {', '.join(sorted(nao_encontrados)[:10])}"
                    + (" ..." if len(nao_encontrados) > 10 else "")
                )

            # Filtrar apenas unidades com itens
            unidades_com_itens = {u: v for u, v in itens_por_unidade.items() if v}

            if not unidades_com_itens:
                st.info("Nenhuma quantidade encontrada na planilha. Preencha as colunas dos estabelecimentos e faça o upload novamente.")
            else:
                st.markdown(f"**Prévia — {len(unidades_com_itens)} solicitação(ões) a criar:**")

                # Mapa produto_id → descricao para exibição
                prod_desc_map = {}
                if not df_produtos.empty:
                    for _, p in df_produtos.iterrows():
                        prod_desc_map[p["id"]] = str(p.get("descricao", p["id"]))

                for unid, itens in unidades_com_itens.items():
                    with st.expander(f"**{unid}** — {len(itens)} produto(s)", expanded=False):
                        rows_prev = [
                            {"Produto": prod_desc_map.get(pid, str(pid)), "Quantidade": qtd}
                            for pid, qtd in itens.items()
                        ]
                        st.dataframe(pd.DataFrame(rows_prev), use_container_width=True, hide_index=True)

                st.markdown("---")
                if st.button("✅ Confirmar e criar solicitações", type="primary", use_container_width=True):
                    try:
                        df_pedidos_reload = ler_df("pedidos")
                        df_itens_reload   = ler_df("itens_pedido")
                        novo_id  = _next_id(df_pedidos_reload)
                        item_id  = _next_id(df_itens_reload)
                        criados  = []

                        for unid, itens in unidades_com_itens.items():
                            append_linha("pedidos", [
                                novo_id, unid, "aberto",
                                usuario["nome"], datetime.datetime.now().isoformat(), "", "",
                            ])
                            for prod_id, qtd in itens.items():
                                append_linha("itens_pedido", [item_id, novo_id, prod_id, qtd])
                                item_id += 1
                            criados.append(f"#{novo_id} — {unid}")
                            novo_id += 1

                        st.cache_data.clear()
                        st.session_state["import_ok"] = (
                            f"✅ {len(criados)} solicitação(ões) criadas com sucesso: "
                            + ", ".join(criados)
                        )
                        st.session_state["upload_v"] = st.session_state.get("upload_v", 0) + 1
                        st.rerun()
                    except Exception as e:
                        st.error(f"Erro ao criar solicitações: {e}")
